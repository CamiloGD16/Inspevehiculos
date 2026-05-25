from docx import Document
from docx.shared import Inches, Pt
import json
import os

# Crear documento
doc = Document()
doc.add_heading('Reporte de Inspección', 0)

# Agregar tabla
table = doc.add_table(rows=1, cols=3)
table.style = 'Light Grid Accent 1'

# Encabezados
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Item'
hdr_cells[1].text = 'Cumple'
hdr_cells[2].text = 'Observación'

# Datos de ejemplo
datos = [
    {'Item': 'Techos', 'Cumple': 'No cumple', 'Observacion': 'Sucio'},
    {'Item': 'Pisos', 'Cumple': 'Cumple', 'Observacion': ''}
]

for dato in datos:
    row_cells = table.add_row().cells
    row_cells[0].text = dato['Item']
    row_cells[1].text = dato['Cumple']
    row_cells[2].text = dato['Observacion']

# Guardar
doc.save('reporte.docx')
print("Archivo creado: reporte.docx")
